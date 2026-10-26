from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Declaracao_Zona_Tecnica_editavel.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.1)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

def set_font(run, size=11, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def add_field(paragraph, label):
    run = paragraph.add_run(f"[{label}]")
    set_font(run, size=11, bold=True, color="1F4E79")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "EAF2F8")
    run._element.get_or_add_rPr().append(shading)
    return run

# Identificação visual simples, sem copiar o logótipo rasterizado.
brand = doc.add_paragraph()
brand.paragraph_format.space_after = Pt(0)
r = brand.add_run("ZonaTécnica")
set_font(r, size=15, bold=True, italic=True, color="111111")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
r = subtitle.add_run("FUTSAL PORTUGAL")
set_font(r, size=7.5, bold=True, color="D96C1F")

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(4)
title.paragraph_format.space_after = Pt(28)
r = title.add_run("DECLARAÇÃO")
set_font(r, size=14, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.first_line_indent = Cm(1.0)
p.paragraph_format.space_after = Pt(0)
add_field(p, "DENOMINAÇÃO DA EMPRESA")
r = p.add_run(", com sede em ")
set_font(r)
add_field(p, "SEDE / MORADA DA EMPRESA")
r = p.add_run(", declara, nos termos e para os efeitos do disposto no artigo 15.º do “Estatuto do Jornalista”, Lei n.º 1/99, de 13 de janeiro, com a redação introduzida pela Lei n.º 64/2007, de 6 de novembro, que ")
set_font(r)
add_field(p, "NOME COMPLETO DO DIRETOR DO SETOR INFORMATIVO")
r = p.add_run(", portador do documento de identificação n.º ")
set_font(r)
add_field(p, "NÚMERO DO DOCUMENTO")
r = p.add_run(", emitido em ")
set_font(r)
add_field(p, "DATA DE EMISSÃO")
r = p.add_run(" por ")
set_font(r)
add_field(p, "ENTIDADE EMISSORA")
r = p.add_run(", desempenha, de forma efetiva e permanente, as funções de Diretor(a) da publicação ")
set_font(r)
add_field(p, "NOME DA PUBLICAÇÃO")
r = p.add_run(", desde ")
set_font(r)
add_field(p, "DATA DE INÍCIO")
r = p.add_run(".")
set_font(r)

doc.add_paragraph().paragraph_format.space_after = Pt(14)

date_p = doc.add_paragraph()
date_p.paragraph_format.space_after = Pt(30)
add_field(date_p, "LOCAL")
r = date_p.add_run(", ")
set_font(r)
add_field(date_p, "DATA")

sig = doc.add_paragraph()
sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
sig.paragraph_format.space_after = Pt(2)
r = sig.add_run("________________________________________")
set_font(r, size=11)

sig_name = doc.add_paragraph()
sig_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
sig_name.paragraph_format.space_after = Pt(2)
add_field(sig_name, "NOME DO REPRESENTANTE DA EMPRESA")

sig_role = doc.add_paragraph()
sig_role.alignment = WD_ALIGN_PARAGRAPH.CENTER
sig_role.paragraph_format.space_after = Pt(0)
r = sig_role.add_run("(assinatura do representante da empresa, com aposição do carimbo)")
set_font(r, size=9, italic=True, color="555555")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Modelo editável — substituir os campos destacados antes de assinar")
set_font(r, size=8, italic=True, color="777777")

doc.core_properties.title = "Declaração - Zona Técnica"
doc.core_properties.subject = "Modelo editável de declaração"
doc.core_properties.author = "Zona Técnica Futsal Portugal"
doc.save(DOCX)
print(DOCX)
