from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from pathlib import Path


OUT = Path(__file__).resolve().parent
DOCX = OUT / "CV_Rui_Simoes.docx"

NAVY = "17324D"
TEAL = "147D82"
INK = "25313B"
MUTED = "60717F"
LIGHT = "EAF2F3"
PALE = "F5F8F9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="D7E1E5", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_font(run, size=None, bold=None, color=None, italic=None, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_text(p, text, **kwargs):
    r = p.add_run(text)
    set_font(r, **kwargs)
    return r


def add_bottom_border(paragraph, color=TEAL, size="16", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def section_title(doc, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Section Title"]
    add_text(p, title.upper(), size=10.5, bold=True, color=NAVY)
    add_bottom_border(p, TEAL, "8", "3")
    return p


def bullet(doc, text, compact=True):
    p = doc.add_paragraph(style="CV Bullet")
    p.paragraph_format.space_after = Pt(1.5 if compact else 3)
    add_text(p, text, size=9.1, color=INK)
    return p


def role(doc, title, org, dates, location=None, bullets=None):
    p = doc.add_paragraph(style="Role Header")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.48), WD_TAB_ALIGNMENT.RIGHT)
    add_text(p, title, size=10.2, bold=True, color=NAVY)
    add_text(p, f"\t{dates}", size=9, bold=True, color=TEAL)
    q = doc.add_paragraph(style="Role Meta")
    q.paragraph_format.keep_with_next = True
    add_text(q, org, size=9.2, bold=True, color=MUTED)
    if location:
        add_text(q, f"  |  {location}", size=9, color=MUTED)
    for item in bullets or []:
        bullet(doc, item)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)
sec.header_distance = Inches(0.25)
sec.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05

for name, before, after in [
    ("Section Title", 7, 4),
    ("Role Header", 4, 0),
    ("Role Meta", 0, 1.5),
]:
    st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.0

bs = styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
bs.font.name = "Aptos"
bs._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
bs._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
bs.paragraph_format.left_indent = Inches(0.18)
bs.paragraph_format.first_line_indent = Inches(-0.13)
bs.paragraph_format.line_spacing = 1.02

# Header
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, "RUI SIMÕES", size=25, bold=True, color=NAVY)
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(5)
add_text(
    p2,
    "ADMINISTRAÇÃO  •  OPERAÇÕES  •  LOGÍSTICA  •  COORDENAÇÃO",
    size=9.4,
    bold=True,
    color=TEAL,
)
add_bottom_border(p2, TEAL, "18", "6")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_before = Pt(2)
contact.paragraph_format.space_after = Pt(5)
add_text(contact, "Azambuja, Portugal  |  +351 912 654 164  |  rui16simoes@gmail.com", size=9, color=MUTED)
add_text(contact, "\nlinkedin.com/in/rui-simões-ab0586b7", size=8.7, color=MUTED)

section_title(doc, "Perfil profissional")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
add_text(
    p,
    "Profissional com mais de 20 anos de experiência em ambientes operacionais e de serviço, "
    "com percurso consolidado em apoio administrativo à produção, logística, gestão de stocks, "
    "transportes, coordenação de equipas e atendimento. Experiência em projetos ferroviários, "
    "controlo de materiais em SAP, compras, planeamento de recursos e articulação entre diferentes "
    "intervenientes. Perfil organizado, responsável e versátil, com forte capacidade de liderança, "
    "comunicação, resolução de problemas e adaptação a contextos exigentes.",
    size=9.3,
    color=INK,
)

section_title(doc, "Competências-chave")
t = doc.add_table(rows=3, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [Inches(2.34), Inches(2.34), Inches(2.34)]
skills = [
    "Apoio administrativo à produção",
    "Gestão operacional e planeamento",
    "Logística, stocks e aprovisionamento",
    "SAP e controlo de materiais",
    "Compras e consulta ao mercado",
    "Coordenação de equipas e turnos",
    "Gestão documental e relatórios",
    "Transportes e expedição",
    "Comunicação e atendimento",
]
for i, row in enumerate(t.rows):
    for j, cell in enumerate(row.cells):
        cell.width = widths[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 70, 100, 70, 100)
        set_cell_shading(cell, PALE if i % 2 == 0 else WHITE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, skills[i * 3 + j], size=8.5, bold=True, color=NAVY)
set_table_borders(t, "D8E4E6", "3")

section_title(doc, "Experiência profissional")
role(
    doc,
    "Técnico Administrativo de Produção",
    "DTE - Instalações Especiais",
    "out 2025 - atual",
    bullets=[
        "Apoio administrativo e operacional à Direção de Produção em projetos ferroviários, assegurando a organização de processos e o suporte às operações em obra.",
        "Coordenação administrativa entre equipas de obra, fiscalização e dono de obra; controlo de assiduidade, entradas e saídas, turnos diurnos e noturnos.",
        "Gestão de compras, consulta ao mercado, alojamentos e mobilização de equipas; registo e controlo de materiais, ferramentas e equipamentos em SAP.",
        "Preparação de documentação técnica e administrativa, apoio a autos de medição, mapas económico-financeiros, relatórios e comunicações internas.",
    ],
)
role(
    doc,
    "Chefe de Estação",
    "Boa Viagem",
    "ago 2015 - out 2025",
    "Arruda dos Vinhos",
    bullets=[
        "Coordenação de movimentos de autocarros, equipas e recursos, garantindo a execução do serviço e o cumprimento dos horários.",
        "Colaboração na elaboração de horários, gestão de substituições de pessoal e viaturas e mobilização de veículos suplementares em períodos de maior afluência.",
        "Controlo do serviço diário e de registos tacográficos, análise de atrasos, gestão de ocorrências, elaboração de relatórios e receção de contas de motoristas.",
    ],
)
role(
    doc,
    "Administrativo de Logística e Gestão de Stocks",
    "STEF Portugal",
    "2008 - 2015",
    "Alverca, Azambuja e Póvoa de Santa Iria",
    bullets=[
        "Gestão de stocks, receção, expedição e aprovisionamento; apoio administrativo aos clientes Auchan, GL, Sogenave, Ferrero, Spar e Terras de Alter.",
        "Responsabilidade pelo economato da plataforma de Alverca e apoio aos processos de qualidade, higienização de armazém e serviço ao cliente.",
    ],
)

doc.add_page_break()

section_title(doc, "Experiência profissional — continuação")
role(
    doc,
    "Assistente Comercial",
    "Companhia Geral de Leilões / Leilocar",
    "2007 - 2008",
    bullets=["Apoio comercial e administrativo numa leiloeira de automóveis, com contacto regular com clientes e suporte aos processos de venda."],
)
role(
    doc,
    "Administrativo",
    "Hospital CUF Infante Santo",
    "2006 - 2007",
    bullets=["Marcação telefónica de exames e consultas e atendimento administrativo a utentes."],
)
role(
    doc,
    "Administrativo de Receção",
    "El Corte Inglés Portugal",
    "2005 - 2006",
    "Lisboa",
    bullets=["Receção administrativa e controlo de produtos destinados à loja."],
)

section_title(doc, "Projetos, liderança e comunicação")
role(
    doc,
    "Fundador e Responsável Editorial",
    "Zona Técnica Futsal Portugal",
    "2014 - atual",
    bullets=[
        "Criação e gestão editorial de uma plataforma digital especializada em futsal português, incluindo estratégia de conteúdos, notícias, entrevistas, análises táticas e cobertura de competições.",
        "Gestão do site e plataformas digitais, coordenação do calendário editorial e cobertura internacional de Campeonatos da Europa e UEFA Futsal Champions League.",
    ],
)
role(
    doc,
    "Treinador de Guarda-Redes, Analista e Coordenador",
    "Vários clubes e projetos de futsal",
    "2011 - atual",
    bullets=[
        "Experiência em treino de guarda-redes, observação e análise de jogo, coordenação técnica e apoio à gestão de equipas, incluindo SL Benfica S17/S20, Reguilas Tires, AF Santarém, Futsal Oeiras e UPVN.",
        "Participação no Campeonato do Mundo de Futsal de 2022 e colaboração com a Seleção Nacional de Pessoas com Diabetes.",
    ],
)

section_title(doc, "Educação e formação")
edu = doc.add_table(rows=4, cols=2)
edu.alignment = WD_TABLE_ALIGNMENT.CENTER
edu.autofit = False
edu.columns[0].width = Inches(1.45)
edu.columns[1].width = Inches(5.55)
education = [
    ("2025 - atual", "TeSP em Gestão Administrativa de Recursos Humanos — ISLA Santarém | Em curso desde setembro de 2025"),
    ("2009 - 2010", "Frequência do 1.º ano da licenciatura em Gestão de Empresas — Escola Superior de Gestão de Santarém"),
    ("2002 - 2005", "Curso Técnico de Transportes — IPTrans, Loures | Equivalência ao 12.º ano, nível III | 14 valores"),
    ("1996 - 2001", "9.º ano — Instituto Militar dos Pupilos do Exército"),
]
for i, (date, text) in enumerate(education):
    c1, c2 = edu.rows[i].cells
    for c in (c1, c2):
        set_cell_margins(c, 70, 95, 70, 95)
        set_cell_shading(c, PALE if i % 2 == 0 else WHITE)
    add_text(c1.paragraphs[0], date, size=8.7, bold=True, color=TEAL)
    add_text(c2.paragraphs[0], text, size=8.7, color=INK)
    c1.paragraphs[0].paragraph_format.space_after = Pt(0)
    c2.paragraphs[0].paragraph_format.space_after = Pt(0)
set_table_borders(edu, "D8E4E6", "3")

section_title(doc, "Formação complementar selecionada")
for text in [
    "Curso de Treinador de Futsal Nível II (2023) e formação contínua em metodologia, plano de jogo, microciclo e treino de jovens (2023-2024).",
    "Formação de Tutores de Futsal (2024).",
    "Ações de formação em combate a incêndios (2015 e 2023).",
    "Curso Europeu de Primeiros Socorros (2009).",
    "Formação em atendimento, serviço ao cliente, higienização e qualidade.",
    "Curso de Mediador de Seguros Ligado - Ramo Não Vida (2008).",
]:
    bullet(doc, text)

section_title(doc, "Idiomas e ferramentas")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
add_text(p, "Português", size=9, bold=True, color=NAVY)
add_text(p, " — língua materna   |   ", size=9, color=INK)
add_text(p, "Inglês", size=9, bold=True, color=NAVY)
add_text(p, " — nível básico (A2)", size=9, color=INK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
add_text(p, "Ferramentas e sistemas: ", size=9, bold=True, color=NAVY)
add_text(p, "SAP; ferramentas de gestão documental e operacional; Video Observer; plataformas digitais e gestão de conteúdos.", size=9, color=INK)

# Footer
for section in doc.sections:
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    add_text(fp, "Rui Simões  •  Curriculum Vitae", size=7.5, color=MUTED)

# Keep table rows together where practical.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

doc.core_properties.title = "Curriculum Vitae - Rui Simões"
doc.core_properties.subject = "Curriculum Vitae"
doc.core_properties.author = "Rui Simões"
doc.core_properties.keywords = "administração, operações, logística, transportes, gestão de stocks, SAP"
doc.save(DOCX)
print(DOCX)
