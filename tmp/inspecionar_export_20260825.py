from pathlib import Path
from openpyxl import load_workbook
import json
import re

p = Path(r"C:\Users\dst7973\Desktop\EXPORT_20260825_125313.xlsx")
wb = load_workbook(p, data_only=False, read_only=True)
print("SHEETS", wb.sheetnames)
for ws in wb.worksheets:
    print("\nSHEET", ws.title, "ROWS", ws.max_row, "COLS", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        print(i, repr(row[:min(ws.max_column, 25)]))
        if i >= 15:
            break

data = json.loads(Path("dados_armazem_catenaria.json").read_text(encoding="utf-8"))
atuais = {}
for e in data.get("encomendas", []):
    pedido = str(e.get("pedido", "")).strip()
    m = re.search(r"Item\s+(\d+)", str(e.get("obs", "")), re.I)
    if not m:
        m = re.search(r"-(\d+)$", str(e.get("id", "")))
    if pedido and m:
        atuais[(pedido, str(int(m.group(1))))] = e

ws = wb[wb.sheetnames[0]]
print("\nNOVAS_OU_RECECOES_A_MAIS")
for r in range(2, ws.max_row + 1):
    item = str(ws.cell(r, 2).value or "").strip()
    pedido = str(ws.cell(r, 10).value or "").strip()
    qtd = float(ws.cell(r, 12).value or 0)
    falta = float(ws.cell(r, 17).value or 0)
    recebida = qtd - falta
    atual = atuais.get((pedido, str(int(float(item))) if item else item))
    atual_rec = float(atual.get("qtdRecebida", 0) or 0) if atual else 0
    if (not atual) or recebida > atual_rec + 1e-9:
        vals = {
            "linha": r, "data": str(ws.cell(r, 6).value)[:10],
            "fornecedor": ws.cell(r, 7).value, "sap": ws.cell(r, 8).value,
            "descricao": ws.cell(r, 9).value, "pedido": pedido, "item": item,
            "qtd": qtd, "recebida_excel": recebida, "recebida_programa": atual_rec,
            "nova_entrada": recebida-atual_rec, "por_fornecer": falta,
            "un": ws.cell(r, 13).value, "existe_encomenda": bool(atual),
        }
        print(json.dumps(vals, ensure_ascii=False, default=str))
