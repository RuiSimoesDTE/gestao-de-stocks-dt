from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\dst7973\Desktop\gestao_de_stocks_dte")
OUT = ROOT / "output" / "candidatura_cp"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "CV_Rui_Simoes_CP.docx"

NAVY = RGBColor(23, 54, 77)
BLUE = RGBColor(33, 103, 126)
GRAY = RGBColor(86, 96, 105)
LIGHT = "DCE8ED"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_repeat_table_layout(table, widths_twips):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths_twips)))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths_twips[idx]))
            tcW.set(qn("w:type"), "dxa")


def set_run(run, size=9.6, bold=False, color=None, italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_rule(paragraph, color="217B8F", size="10"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    set_run(r, size=10.3, bold=True, color=BLUE)
    add_rule(p, color="78AAB8", size="6")
    return p


def add_role(doc, title, employer, dates, bullets=None, summary=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(title), size=10.1, bold=True, color=NAVY)
    if employer:
        set_run(p.add_run(f" | {employer}"), size=9.8, bold=True, color=GRAY)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(1.5)
    p2.paragraph_format.keep_with_next = bool(bullets or summary)
    set_run(p2.add_run(dates), size=8.8, color=GRAY, italic=True)
    if summary:
        ps = doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(2)
        ps.paragraph_format.line_spacing = 1.05
        set_run(ps.add_run(summary), size=9.3)
    for item in bullets or []:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.left_indent = Inches(0.2)
        bp.paragraph_format.first_line_indent = Inches(-0.13)
        bp.paragraph_format.space_after = Pt(1.3)
        bp.paragraph_format.line_spacing = 1.02
        set_run(bp.add_run(item), size=9.25)


doc = Document()
sec = doc.sections[0]
# A4 is a named override appropriate to a Portuguese recruitment document.
sec.page_width = Inches(8.2677)
sec.page_height = Inches(11.6929)
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.62)
sec.right_margin = Inches(0.62)
sec.header_distance = Inches(0.22)
sec.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.6)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05
for style_name in ("Heading 1", "Heading 2", "Heading 3"):
    st = styles[style_name]
    st.font.name = "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
list_style = styles["List Bullet"]
list_style.font.name = "Aptos"
list_style.font.size = Pt(9.25)

# Customer-pack-inspired CV masthead: strong left-aligned identity and compact contact row.
name_p = doc.add_paragraph()
name_p.paragraph_format.space_after = Pt(0)
set_run(name_p.add_run("RUI SIMÕES"), size=23, bold=True, color=NAVY)

tag_p = doc.add_paragraph()
tag_p.paragraph_format.space_after = Pt(4)
set_run(tag_p.add_run("ASSISTENTE ADMINISTRATIVO | SAP | GESTÃO DE ASSIDUIDADE | RECURSOS HUMANOS"), size=10.2, bold=True, color=BLUE)

contact = doc.add_table(rows=1, cols=2)
set_repeat_table_layout(contact, [5160, 5160])
no_table_borders(contact)
for c in contact.rows[0].cells:
    set_cell_shading(c, LIGHT)
    set_cell_margins(c, top=85, bottom=85, start=120, end=120)
p = contact.cell(0, 0).paragraphs[0]
p.paragraph_format.space_after = Pt(0)
set_run(p.add_run("Azambuja, Lisboa  |  912 654 164  |  rui16simoes@gmail.com"), size=8.8, color=NAVY)
p = contact.cell(0, 1).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_after = Pt(0)
set_run(p.add_run("linkedin.com/in/rui-simões-ab0586b7"), size=8.8, color=NAVY)

section_heading(doc, "Perfil profissional")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
p.paragraph_format.line_spacing = 1.08
set_run(p.add_run(
    "Profissional administrativo com mais de 15 anos de experiência em suporte administrativo, gestão operacional e coordenação de equipas nos setores dos transportes, logística e projetos ferroviários. Experiência atual em controlo de assiduidade, organização de turnos, gestão documental, elaboração de relatórios e utilização de SAP. A frequentar o TeSP em Gestão Administrativa de Recursos Humanos. Perfil responsável, organizado, rigoroso e habituado a articular com diferentes equipas e interlocutores."
), size=9.45)

section_heading(doc, "Competências-chave")
skills = doc.add_table(rows=2, cols=3)
set_repeat_table_layout(skills, [3440, 3440, 3440])
no_table_borders(skills)
skill_items = [
    "SAP - registo e controlo de materiais",
    "Controlo de assiduidade e turnos",
    "Gestão documental e arquivo",
    "Microsoft Office - nível intermédio",
    "Relatórios e comunicações internas",
    "Coordenação administrativa de equipas",
]
for cell, txt in zip([c for row in skills.rows for c in row.cells], skill_items):
    set_cell_margins(cell, top=45, bottom=45, start=50, end=80)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("• "), size=9, bold=True, color=BLUE)
    set_run(p.add_run(txt), size=9)

section_heading(doc, "Experiência profissional")
add_role(doc, "Técnico Administrativo de Produção", "DTE - Instalações Especiais", "outubro 2025 - atual", bullets=[
    "Apoio administrativo e operacional ao Diretor de Produção em projetos ferroviários.",
    "Controlo de assiduidade, entradas e saídas de trabalhadores e organização de turnos.",
    "Preparação de listagens de equipas para turnos diurnos e noturnos.",
    "Coordenação administrativa entre equipas de obra, fiscalização e dono de obra.",
    "Registo e controlo de materiais, ferramentas e equipamentos em SAP.",
    "Preparação de documentação, relatórios, comunicações internas e mapas de controlo.",
    "Apoio à gestão de compras, alojamentos e mobilização de equipas.",
])

add_role(doc, "Chefe de Estação", "Boa Viagem", "agosto 2015 - outubro 2025", bullets=[
    "Coordenação e orientação do movimento de autocarros e das equipas afetas à operação.",
    "Apoio à elaboração de horários, substituição de pessoal e reforço de meios.",
    "Controlo do serviço diário, cumprimento de horários e análise de atrasos e ocorrências.",
    "Verificação de registos tacográficos, elaboração de relatórios e receção de contas de motoristas.",
])

add_role(doc, "Administrativo", "STEF - Póvoa de Santa Iria", "2013 - 2015", summary="Apoio administrativo dedicado ao cliente Auchan Portugal, no contexto de operações logísticas.")
add_role(doc, "Administrativo de Gestão de Stocks", "STEF", "2011 - 2013", summary="Gestão administrativa de stocks, receção, expedição e aprovisionamento para o cliente GL.")

doc.add_page_break()
section_heading(doc, "Experiência profissional - continuação")
add_role(doc, "Administrativo", "STEF - Alverca", "2008 - 2011", bullets=[
    "Responsável pelo economato da plataforma de Alverca e gestão de stocks.",
    "Apoio administrativo aos clientes Sogenave, Ferrero, Spar e Terras de Alter.",
    "Atividades de qualidade ao cliente, higienização e qualidade do armazém.",
])
add_role(doc, "Assistente Comercial", "Companhia Geral de Leilões / Leilocar", "2007 - 2008", summary="Apoio comercial e administrativo numa leiloeira de automóveis.")
add_role(doc, "Administrativo Hospitalar", "CUF - Hospital Infante Santo", "2006 - 2007", summary="Atendimento telefónico e marcação de consultas e exames.")
add_role(doc, "Administrativo de Receção", "El Corte Inglés Portugal", "2005 - 2006", summary="Receção e registo administrativo de produtos destinados à loja.")

section_heading(doc, "Formação académica")
add_role(doc, "TeSP em Gestão Administrativa de Recursos Humanos", "ISLA Santarém", "2025 - atual", summary="Formação em curso na área da gestão administrativa de pessoas e processos de Recursos Humanos.")
add_role(doc, "Frequência universitária - 1.º ano de Gestão de Empresas", "Escola Superior de Gestão de Santarém", "2009 - 2010")
add_role(doc, "Técnico de Transportes - qualificação profissional de nível III", "", "2002 - 2005", summary="Curso com equivalência ao 12.º ano de escolaridade.")
add_role(doc, "Formação no Instituto Militar dos Pupilos do Exército", "", "1996 - 2001")

section_heading(doc, "Projeto complementar")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_run(p.add_run("Fundador e responsável editorial | Zona Técnica Futsal Portugal | 2014 - atual"), size=9.5, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
set_run(p.add_run("Gestão editorial de uma plataforma digital, produção e revisão de conteúdos, realização de entrevistas e coordenação do calendário de publicações."), size=9.25)

# Quiet footer with page numbering field.
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(2)
set_run(fp.add_run("Página "), size=8, color=GRAY)
run = fp.add_run()
set_run(run, size=8, color=GRAY)
begin = OxmlElement("w:fldChar")
begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
separate = OxmlElement("w:fldChar")
separate.set(qn("w:fldCharType"), "separate")
text_el = OxmlElement("w:t")
text_el.text = "1"
end = OxmlElement("w:fldChar")
end.set(qn("w:fldCharType"), "end")
run._r.extend([begin, instr, separate, text_el, end])
set_run(fp.add_run(" de 2"), size=8, color=GRAY)

doc.core_properties.title = "Curriculum Vitae - Rui Simões"
doc.core_properties.subject = "Candidatura CP - Assistente Administrativo - RE 35DPC/2026"
doc.core_properties.author = "Rui Simões"
doc.core_properties.keywords = "Assistente Administrativo, SAP, Recursos Humanos, assiduidade, CP"
doc.save(DOCX)
print(DOCX)
