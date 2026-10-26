from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict
import json, re

xlsx = Path(r"C:\Users\dst7973\Desktop\EXPORT_20260825_125313.xlsx")
wb = load_workbook(xlsx, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
db = json.loads(Path("dados_armazem_catenaria.json").read_text(encoding="utf-8"))
mats = {str(m.get("id")): m for m in db.get("materiais", [])}

mov_por_pc = defaultdict(lambda: {"entrada": 0.0, "saida": 0.0, "linhas": []})
for mov in db.get("movimentos", []):
    data = str(mov.get("data", ""))
    if not data.startswith("2026-08"):
        continue
    texto = " ".join(str(mov.get(k, "")) for k in ("obs", "doc", "origem", "destino"))
    pcs = re.findall(r"45\d{8}", texto)
    for pc in set(pcs):
        tipo = str(mov.get("tipo", "")).lower()
        if tipo in ("entrada", "saida"):
            mov_por_pc[pc][tipo] += float(mov.get("qtd", 0) or 0)
            mov_por_pc[pc]["linhas"].append(mov.get("id"))

res = defaultdict(lambda: {"recebidas": [], "pendentes": [], "pcs": set()})
for r in range(2, ws.max_row + 1):
    dt = ws.cell(r, 6).value
    if not getattr(dt, "year", None) == 2026 or dt.month != 8:
        continue
    forn = str(ws.cell(r, 7).value or "").strip()
    pedido = str(ws.cell(r, 10).value or "").strip()
    qtd = float(ws.cell(r, 12).value or 0)
    falta = float(ws.cell(r, 17).value or 0)
    rec = qtd - falta
    linha = {
        "linha": r, "data": dt.strftime("%Y-%m-%d"), "pc": pedido,
        "item": str(ws.cell(r, 2).value or ""), "sap": str(ws.cell(r, 8).value or ""),
        "desc": str(ws.cell(r, 9).value or ""), "qtd": qtd, "recebida": rec,
        "falta": falta, "un": str(ws.cell(r, 13).value or ""),
    }
    res[forn]["pcs"].add(pedido)
    (res[forn]["recebidas"] if rec > 0 else res[forn]["pendentes"]).append(linha)

for forn, x in sorted(res.items()):
    if not x["recebidas"]:
        continue
    print("FORNECEDOR", forn)
    print("PCS", sorted(x["pcs"]))
    print("RECEBIDAS_LINHAS", len(x["recebidas"]), "QTD_TOTAL", sum(z["recebida"] for z in x["recebidas"]))
    for z in x["recebidas"]:
        mm = mov_por_pc.get(z["pc"], {})
        print(json.dumps({**z, "mov_entrada_pc": mm.get("entrada", 0), "mov_saida_pc": mm.get("saida", 0)}, ensure_ascii=False))
    print()
